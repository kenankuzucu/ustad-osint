# -*- coding: utf-8 -*-
"""
ÜSTAD OSINT — EK2C MODÜLLERİ (v1.6)
Saldırı Akışı (SANS ISC) · Sır Avcısı · Yerel Ağ · Şifre Kontrol (HIBP) ·
Rapor Paketi (Word/HTML) · Bildirim Merkezi

Ana çekirdek (osint-araclar.py / osint-ek.py) "/api/arac/<ad>" yollarını buraya devreder.
Kendi küçük yardımcılarını taşır; BAŞKA proje dosyalarına dokunmaz.
Tek başına denemek için:
    python osint-ek2c.py saldiriakisi
    python osint-ek2c.py yerelag
"""
import hashlib
import html as _html
import json
import math
import os
import re
import socket
import ssl
import subprocess
import time
import urllib.error
import urllib.parse
import urllib.request

KOK = os.path.dirname(os.path.abspath(__file__))
VERI = os.path.join(KOK, "veri")
ONBELLEK = os.path.join(VERI, "ek2-onbellek.json")          # 15 dk önbellek dosyası
YERELAG_HAFIZA = os.path.join(VERI, "yerelag.json")         # Yerel ağ cihaz hafızası
BILDIRIM_DOSYA = os.path.join(VERI, "bildirim.json")        # Bildirim ayar + olay günlüğü
ENV_DOSYA = os.path.join(os.path.expanduser("~"), "AppData", "Local", "hermes", ".env")
RAPOR_KLASOR = os.path.join(
    os.path.expanduser("~"), "OneDrive", "Desktop", "USTAD-OSINT-RAPORLAR")

# Bazı ortamlarda ('+' gerekli) — izinli/kendi sistem bağlamı
UA = "UstadOSINT/1.6 (+kendi sistem, izinli)"
SURUM = "1.6"

# ------------------------------------------------------------------ SSL BAĞLAMI
def _ssl_baglam():
    """Önce certifi (gerçek sertifika doğrulaması), olmazsa sistem CA deposu."""
    try:
        import certifi
        return ssl.create_default_context(cafile=certifi.where())
    except Exception:
        try:
            return ssl.create_default_context()
        except Exception:
            return None


BAGLAM = _ssl_baglam()
# Yedek: Kaspersky MITM gibi araya giren kurulumlarda doğrulamasız bağlam (yalnız yedek yol)
try:
    BAGLAM_GUVENSIZ = ssl._create_unverified_context()
except Exception:
    BAGLAM_GUVENSIZ = None


def _cek(url, zaman=15, baslik=None, veri=None, tip="GET", metin=True):
    """Basit HTTP istemcisi. (durum_kodu, gövde) döner; ASLA hata fırlatmaz.

    Önce doğrulanmış bağlam (certifi), başarısız olursa doğrulamasız yedek denenir.
    """
    istek = urllib.request.Request(url, data=veri)
    istek.add_header("User-Agent", UA)
    istek.add_header("Accept", "application/json,text/html,text/plain,*/*")
    if veri is not None:
        istek.add_header("Content-Type", "application/json")
    for k, v in (baslik or {}).items():
        istek.add_header(k, v)
    if tip != "GET":
        istek.get_method = lambda: tip
    for baglam in (BAGLAM, BAGLAM_GUVENSIZ):
        try:
            with urllib.request.urlopen(istek, timeout=zaman, context=baglam) as y:
                ham = y.read(2000000)
                return y.status, (ham.decode("utf-8", "replace") if metin else ham)
        except urllib.error.HTTPError as e:
            try:
                ham = e.read(400000)
                return e.code, (ham.decode("utf-8", "replace") if metin else ham)
            except Exception:
                return e.code, ("" if metin else b"")
        except Exception:
            continue
    return 0, ("" if metin else b"")


def _json_cek(url, zaman=15, baslik=None, veri=None, tip="GET"):
    """(durum, python_nesnesi, ham_gövde) döner."""
    kod, govde = _cek(url, zaman, baslik, veri, tip)
    try:
        return kod, json.loads(govde), govde
    except Exception:
        return kod, None, govde


# ------------------------------------------------------------------- ÖNBELLEK
def _onbellek_oku():
    try:
        with open(ONBELLEK, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def _onbellek_yaz(d):
    try:
        os.makedirs(VERI, exist_ok=True)
        with open(ONBELLEK, "w", encoding="utf-8") as f:
            json.dump(d, f, ensure_ascii=False)
    except Exception:
        pass


def _onbellekli(anahtar, saniye, uret, zorla=False):
    """`zorla` ise önbellek atlanır; aksi halde saniye kadar geçerli kayıt döner."""
    d = _onbellek_oku()
    if not zorla:
        kayit = d.get(anahtar)
        if kayit and time.time() - kayit.get("_zaman", 0) < saniye:
            v = dict(kayit)
            v["onbellek"] = True
            return v
    v = uret()
    if isinstance(v, dict) and "hata" not in v:
        kayda = dict(v)
        kayda["_zaman"] = time.time()
        d[anahtar] = kayda
        _onbellek_yaz(d)
    return v


def _zorla_mi(sorgu, govde):
    for kaynak in (sorgu or {}, govde or {}):
        if str(kaynak.get("zorla") or "").lower() in ("1", "true", "evet", "yes"):
            return True
    return False


# --------------------------------------------------------------- ortak yardımcı
def _kisalt(m, bas=4, son=4):
    """Bir sırrı maskele: yalnız ilk `bas` + son `son` karakter + '…'."""
    s = str(m or "")
    if len(s) <= bas + son:
        return s[:2] + "…" + s[-2:] if len(s) > 4 else "…"
    return s[:bas] + "…" + s[-son:]


def _entropi(s):
    """Shannon entropisi (bit/karakter)."""
    if not s:
        return 0.0
    s = str(s)
    sayac = {}
    for c in s:
        sayac[c] = sayac.get(c, 0) + 1
    n = len(s)
    h = 0.0
    for c, k in sayac.items():
        p = k / n
        h -= p * math.log2(p)
    return h


def _github_token():
    """GITHUB_TOKEN'ı ortam değişkeninden ya da hermes .env dosyasından okur.

    DİKKAT: değer hiçbir zaman çıktıya/rapora/log'a yazılmaz.
    """
    tok = os.environ.get("GITHUB_TOKEN") or os.environ.get("GH_TOKEN")
    if tok:
        return tok.strip()
    try:
        with open(ENV_DOSYA, "r", encoding="utf-8") as f:
            for satir in f:
                satir = satir.strip()
                if satir.startswith("GITHUB_TOKEN") and "=" in satir:
                    return satir.split("=", 1)[1].strip().strip('"').strip("'")
    except Exception:
        pass
    return ""


# =================================================================== 7) SALDIRI AKIŞI
def saldiriakisi(zorla=False):
    """SANS ISC DShield — dünyada en çok saldıran IP'ler + konum (ip-api, yedek ipwho.is)."""
    anahtar = "saldiriakisi::topips"

    def uret():
        kod, j, _ = _json_cek("https://isc.sans.edu/api/topips/records/20?json", 20)
        if kod != 200 or not isinstance(j, list):
            return {"hata": "SANS ISC verisi alınamadı (durum %s). Ağ/engel olabilir." % kod}
        ham = []
        for r in j[:20]:
            try:
                ham.append({
                    "sira": int(r.get("rank") or len(ham) + 1),
                    "ip": str(r.get("source") or "").strip(),
                    "rapor": int(r.get("reports") or 0),
                    "hedef": int(r.get("targets") or 0),
                })
            except Exception:
                continue
        ipler = [x["ip"] for x in ham if x["ip"]]
        konum = _konum_toplu(ipler)                # {ip: {ulke, sehir, enlem, boylam}}
        kayitlar = []
        for x in ham:
            k = konum.get(x["ip"]) or {}
            kayitlar.append({
                "sira": x["sira"], "ip": x["ip"], "rapor": x["rapor"], "hedef": x["hedef"],
                "ulke": k.get("ulke") or "", "sehir": k.get("sehir") or "",
                "enlem": k.get("enlem"), "boylam": k.get("boylam"),
            })
        toplam = sum(x["rapor"] for x in kayitlar)
        return {
            "kaynak": "SANS ISC DShield",
            "kayitlar": kayitlar,
            "toplam_rapor": toplam,
            "konum_kaynagi": konum.get("_kaynak", ""),
            "guncelleme": time.strftime("%Y-%m-%d %H:%M:%S"),
            "not": "İzinli/kendi sistem bağlamında pasif istihbarat kaynağıdır; hedeflere dokunulmaz.",
        }

    return _onbellekli(anahtar, 900, uret, zorla)   # 15 dakika


def _konum_toplu(ipler):
    """ip-api batch (tek istek) → başarısızsa en fazla 20 IP için ipwho.is tek tek."""
    sonuc = {}
    if not ipler:
        return sonuc
    v = json.dumps([{"query": ip} for ip in ipler[:100]]).encode("utf-8")
    kod, j, _ = _json_cek(
        "http://ip-api.com/batch?fields=status,country,city,lat,lon,query", 20, None, v, "POST")
    if kod == 200 and isinstance(j, list):
        for r in j:
            if not isinstance(r, dict) or r.get("status") != "success":
                continue
            sonuc[str(r.get("query"))] = {
                "ulke": r.get("country") or "", "sehir": r.get("city") or "",
                "enlem": r.get("lat"), "boylam": r.get("lon"),
            }
        if sonuc:
            sonuc["_kaynak"] = "ip-api.com"
            return sonuc
    # ---- yedek: ipwho.is
    for ip in ipler[:20]:
        kod, j, _ = _json_cek("https://ipwho.is/" + urllib.parse.quote(ip), 12)
        if kod == 200 and isinstance(j, dict) and j.get("success"):
            sonuc[ip] = {
                "ulke": j.get("country") or "", "sehir": j.get("city") or "",
                "enlem": j.get("latitude"), "boylam": j.get("longitude"),
            }
    if sonuc:
        sonuc["_kaynak"] = "ipwho.is"
    return sonuc


# ==================================================================== 8) SIR AVCISI
SIR_DESENLER = [
    # (ad, derleme, önem, öneri, sır_grubu)  — sır_grubu: maskelenecek yakalama grubu (0=tüm eşleşme)
    ("AWS Erişim Anahtarı", re.compile(r"AKIA[0-9A-Z]{16}"), "kritik",
     "anahtarı iptal et ve ortam değişkenine taşı", 0),
    ("GitHub token", re.compile(r"(ghp_|gho_|ghu_|ghs_|ghr_)[A-Za-z0-9]{16,}|github_pat_[A-Za-z0-9_]{20,}"),
     "kritik", "anahtarı iptal et ve ortam değişkenine taşı", 0),
    ("Google API anahtarı", re.compile(r"AIza[0-9A-Za-z\-_]{35}"), "yüksek",
     "API anahtarını sınırla (HTTP referer/IP) ve yenile", 0),
    ("Slack token", re.compile(r"xox[baprs]-[A-Za-z0-9-]{10,}"), "kritik",
     "Slack token'ı iptal et ve rotasyon yap", 0),
    ("Stripe canlı anahtar", re.compile(r"sk_live_[A-Za-z0-9]{16,}"), "kritik",
     "Stripe anahtarını iptal et (dashboard → API keys)", 0),
    ("OpenAI anahtar", re.compile(r"sk-[A-Za-z0-9\-_]{20,}"), "yüksek",
     "OpenAI anahtarını iptal et ve kullanım limiti koy", 0),
    ("Telegram bot token", re.compile(r"\b\d{8,10}:[A-Za-z0-9_\-]{35}\b"), "kritik",
     "BotFather → /revoke ile token yenile", 0),
    ("Özel anahtar", re.compile(r"-----BEGIN (?:RSA|OPENSSH|EC|PGP|DSA) PRIVATE KEY-----"), "kritik",
     "özel anahtarı dosyadan çıkar, izinleri 600 yap, yeni anahtar üret", 0),
    ("JWT", re.compile(r"\beyJ[A-Za-z0-9_\-]{10,}\.[A-Za-z0-9_\-]{10,}\.[A-Za-z0-9_\-]{10,}\b"), "yüksek",
     "JWT imza sırrını değiştir; token'ları kısa ömürlü yap", 0),
    ("Bağlantı dizesi", re.compile(r"(?:mysql|postgres|postgresql|mongodb|redis|amqp)://[^\s'\"<>]{8,}"),
     "kritik", "bağlantı dizesindeki parolayı döndür; parolayı ortam değişkenine taşı", 0),
    # Genel kalıp: entropi >= 3.0 şartı ile değerlendirilir
    ("Genel gizli değer", re.compile(
        r"(?i)\b(password|passwd|pwd|api_key|apikey|secret|token|parola)\b\s*[:=]\s*[\"']([^\"']{8,})[\"']"),
     "yüksek", "sırrı koddan çıkar; ortam değişkeni/şifre kasası kullan", 2),
]
SIR_UZANTI = {".py", ".js", ".json", ".txt", ".md", ".env", ".ini", ".cfg", ".yml", ".yaml",
              ".java", ".html", ".htm", ".ps1", ".sh", ".xml", ".properties", ".sql"}
SIR_ATLA_KLASOR = {"node_modules", ".git", "__pycache__", "venv", ".venv", "build",
                   "dist", "tools", "appdata"}
SIR_MAX_BAYT = 2 * 1024 * 1024


def siravc(yol="", limit=400, github=False, zorla=False):
    """Yerel sır avı (ağ kullanmaz). ASLA tam sır değeri döndürmez."""
    yol = str(yol or "").strip() or KOK
    try:
        limit = int(limit or 400)
    except Exception:
        limit = 400
    limit = max(1, min(limit, 5000))
    if not os.path.isdir(yol):
        return {"hata": "klasör bulunamadı: " + yol}
    baslama = time.time()
    bulgular, taranan, atlanan = [], 0, 0
    for kok_yol, klasorler, dosyalar in os.walk(yol):
        # atlanacak klasörleri buda (yerinde değiştir)
        klasorler[:] = [k for k in klasorler
                        if k.lower() not in SIR_ATLA_KLASOR and not k.startswith(".")]
        for ad in dosyalar:
            uzanti = os.path.splitext(ad)[1].lower()
            tam = os.path.join(kok_yol, ad)
            if uzanti not in SIR_UZANTI:
                atlanan += 1
                continue
            try:
                if os.path.getsize(tam) > SIR_MAX_BAYT:
                    atlanan += 1
                    continue
            except Exception:
                atlanan += 1
                continue
            if taranan >= limit:
                atlanan += 1
                continue
            taranan += 1
            try:
                with open(tam, "r", encoding="utf-8", errors="replace") as f:
                    for no, satir in enumerate(f, 1):
                        if len(satir) > 4000:            # aşırı uzun satırları kırp
                            satir = satir[:4000]
                        for ad2, der, onem, oneri, grup in SIR_DESENLER:
                            m = der.search(satir)
                            if not m:
                                continue
                            try:
                                # m.groups() DEMET döner; sayı için len() şart — aksi halde
                                # "tuple >= int" TypeError fırlatır ve tüm eşleşme maskeye düşer.
                                ham = (m.group(grup) if grup <= len(m.groups()) else m.group(0))
                            except Exception:
                                ham = m.group(0)
                            # genel kalıp için entropi eşiği
                            if ad2 == "Genel gizli değer" and _entropi(ham) < 3.0:
                                continue
                            bulgular.append({
                                "dosya": tam, "satir": no, "tur": ad2, "onem": onem,
                                "maske": _kisalt(ham), "oneri": oneri,
                            })
            except Exception:
                atlanan += 1
                continue
    kritik = sum(1 for b in bulgular if b["onem"] == "kritik")
    yuksek = sum(1 for b in bulgular if b["onem"] == "yüksek")
    sure = round(time.time() - baslama, 2)
    sonuc = {
        "kok": yol, "taranan_dosya": taranan, "atlanan": atlanan, "sure_sn": sure,
        "bulgular": bulgular[:500], "kritik": kritik, "yuksek": yuksek,
        "toplam_bulgu": len(bulgular),
        "ozet": ("%d dosya tarandı, %d bulgu (%d kritik, %d yüksek) — %.2f sn"
                 % (taranan, len(bulgular), kritik, yuksek, sure)),
        "kaynak": "yerel dosya sistemi (ağ yok)",
        "not": "Bulgularda sır değerleri maskelenmiştir (yalnız ilk 4 + son 4 karakter).",
    }
    if github:
        sonuc["github"] = _github_kod_arama(zorla)
    return sonuc


def _github_kod_arama(zorla=False):
    """github=1: token varsa GitHub kod arama API'si; reddedilirse dürüstçe bildirir."""
    tok = _github_token()
    if not tok:
        return {"durum": "GitHub token bulunamadı — yalnız yerel tarama yapıldı"}
    bas = {"Authorization": "token " + tok, "Accept": "application/vnd.github+json"}
    kod, j, _ = _json_cek("https://api.github.com/user", 15, bas)
    login = (j or {}).get("login") if isinstance(j, dict) else ""
    if kod != 200 or not login:
        return {"durum": "GitHub token doğrulanamadı (durum %s) — yalnız yerel tarama yapıldı" % kod}
    sonuclar, reddedildi = [], False
    for desen in ("AKIA", "ghp_", "sk_live_", "xoxb-"):
        q = urllib.parse.quote("%s user:%s" % (desen, login))
        kod, j, _ = _json_cek("https://api.github.com/search/code?q=" + q, 20, bas)
        if kod != 200 or not isinstance(j, dict):
            reddedildi = True
            break
        for o in (j.get("items") or [])[:5]:
            sonuclar.append({"desen": desen, "depo": (o.get("repository") or {}).get("full_name"),
                             "yol": o.get("path"), "url": o.get("html_url")})
    if reddedildi:
        return {"durum": "kod arama API'si bu token ile reddedildi — yalnız yerel tarama yapıldı"}
    return {"durum": "tamam", "kullanici": login, "bulunan": len(sonuclar), "kayitlar": sonuclar}


# ====================================================================== 9) YEREL AĞ
# MAC üretici (OUI) tablosu — ön ek ':'.join ile eşleşir; bulunamazsa 'bilinmiyor'
MAC_URETICI = {
    # Sanallaştırma
    "00:50:56": "VMware", "00:0c:29": "VMware", "00:05:69": "VMware", "00:1c:14": "VMware",
    "08:00:27": "VirtualBox", "0a:00:27": "VirtualBox", "52:54:00": "QEMU/KVM",
    # Ağ donanımı
    "3c:52:a1": "TP-Link", "50:c7:bf": "TP-Link", "00:1d:0f": "TP-Link", "a4:2b:b0": "TP-Link",
    "ec:08:6b": "TP-Link", "b0:48:7a": "TP-Link", "f4:f2:6d": "TP-Link", "18:a6:f7": "TP-Link",
    "c0:25:e9": "TP-Link", "60:32:b1": "TP-Link", "14:cc:20": "TP-Link", "98:da:c4": "TP-Link",
    "ac:84:c6": "TP-Link", "90:9a:4a": "TP-Link",
    "00:1e:58": "D-Link", "00:24:01": "D-Link", "1c:7e:e5": "D-Link", "14:d6:4d": "D-Link",
    "00:18:e7": "Cisco", "00:1b:d4": "Cisco", "00:24:97": "Cisco", "00:26:0b": "Cisco",
    "00:1a:a1": "Cisco", "00:0f:b5": "Netgear", "20:e5:2a": "Netgear", "a0:40:a0": "Netgear",
    "00:14:6c": "Netgear", "00:1f:33": "Netgear", "44:94:fc": "Netgear",
    "00:1d:7e": "Cisco-Linksys", "48:f8:b3": "Cisco-Linksys", "00:22:6b": "Cisco-Linksys",
    "00:1e:e5": "Cisco-Linksys", "58:6d:8f": "Cisco-Linksys", "00:0c:41": "Linksys",
    "00:13:49": "Zyxel", "5c:f4:ab": "Zyxel", "b0:b2:dc": "Zyxel",
    "c8:3a:35": "Tenda", "00:b0:0c": "Tenda", "08:10:78": "Tenda", "5c:f9:38": "Tenda",
    "00:0c:42": "MikroTik", "4c:5e:0c": "MikroTik", "6c:3b:6b": "MikroTik", "74:4d:28": "MikroTik",
    "dc:2c:6e": "MikroTik", "e4:8d:8c": "MikroTik",
    "00:04:0e": "AVM (Fritz!Box)", "08:96:d7": "AVM (Fritz!Box)", "34:31:c4": "AVM (Fritz!Box)",
    "38:2c:4a": "AVM (Fritz!Box)",
    "f0:9f:c2": "Ubiquiti", "24:a4:3c": "Ubiquiti", "78:8a:20": "Ubiquiti", "04:18:d6": "Ubiquiti",
    "44:d9:e7": "Ubiquiti", "00:15:6d": "Ubiquiti", "74:ac:b9": "Ubiquiti", "80:2a:a8": "Ubiquiti",
    "00:0b:86": "Aruba/HPE", "6c:f3:7f": "Aruba/HPE", "84:d4:7e": "Aruba/HPE",
    # Apple
    "00:1c:b3": "Apple", "00:25:00": "Apple", "f0:18:98": "Apple", "a4:83:e7": "Apple",
    "3c:07:54": "Apple", "00:1b:63": "Apple", "00:1e:c2": "Apple", "7c:d1:c3": "Apple",
    "68:ab:1e": "Apple", "00:17:f2": "Apple", "00:23:df": "Apple", "34:c0:59": "Apple",
    "60:fb:42": "Apple", "00:0d:93": "Apple", "00:16:cb": "Apple", "9c:04:eb": "Apple",
    "00:26:bb": "Apple", "28:cf:e9": "Apple", "88:66:a5": "Apple", "b8:17:c2": "Apple",
    "dc:2b:2a": "Apple", "18:65:90": "Apple", "00:1f:f3": "Apple", "40:6c:8f": "Apple",
    "c8:2a:14": "Apple", "00:25:bc": "Apple",
    # Tek kart üreticileri
    "b8:27:eb": "Raspberry Pi", "dc:a6:32": "Raspberry Pi", "e4:5f:01": "Raspberry Pi",
    "28:cd:c1": "Raspberry Pi",
    "00:e0:4c": "Realtek", "52:54:ab": "Realtek", "00:1b:21": "Intel", "00:15:17": "Intel",
    "3c:97:0e": "Intel", "8c:16:45": "Intel",
    # Telefon / TV / IoT
    "00:e0:fc": "Huawei", "28:3c:e4": "Huawei", "48:db:50": "Huawei", "70:72:3c": "Huawei",
    "64:b4:73": "Xiaomi", "78:11:dc": "Xiaomi", "8c:be:be": "Xiaomi", "f8:a4:5f": "Xiaomi",
    "50:8f:4c": "Xiaomi", "28:6c:07": "Xiaomi",
    "00:12:fb": "Samsung", "00:15:99": "Samsung", "00:1d:25": "Samsung", "50:32:75": "Samsung",
    "8c:77:12": "Samsung", "e8:50:8b": "Samsung", "5c:0a:5b": "Samsung", "84:25:db": "Samsung",
    "00:1c:62": "LG", "a8:16:b2": "LG", "c4:36:6c": "LG",
    "00:04:1f": "Sony", "00:13:a9": "Sony", "fc:0f:e6": "Sony",
    "00:0d:4b": "Roku", "ac:3a:7a": "Roku", "b0:a7:37": "Roku", "cc:6d:a0": "Roku",
    "00:fc:8b": "Amazon", "40:b4:cd": "Amazon", "44:65:0d": "Amazon", "68:37:e9": "Amazon",
    "74:c2:46": "Amazon", "a0:02:dc": "Amazon", "f0:27:2d": "Amazon", "fc:65:de": "Amazon",
    "00:1a:11": "Google/Nest", "3c:5a:b4": "Google/Nest", "48:d6:d5": "Google/Nest",
    "64:16:66": "Google/Nest", "f4:f5:d8": "Google/Nest", "6c:ad:f8": "Google/Nest",
    "24:0a:c4": "Espressif (ESP)", "5c:cf:7f": "Espressif (ESP)", "3c:71:bf": "Espressif (ESP)",
    "84:cc:a8": "Espressif (ESP)", "a4:cf:12": "Espressif (ESP)", "cc:50:e3": "Espressif (ESP)",
    "dc:4f:22": "Espressif (ESP)", "24:6f:28": "Espressif (ESP)", "7c:9e:bd": "Espressif (ESP)",
    "48:3f:da": "Espressif (ESP)", "8c:aa:b5": "Espressif (ESP)", "c4:4f:33": "Espressif (ESP)",
    "10:52:1c": "Tuya", "68:57:2d": "Tuya", "84:0d:8e": "Tuya/Sonoff", "d8:f1:5b": "Tuya",
    "34:ea:34": "Sonoff/ITEAD", "60:01:94": "Sonoff/ITEAD", "3c:61:05": "Sonoff/ITEAD",
    "24:62:ab": "Sonoff/ITEAD",
    "00:17:88": "Philips Hue", "ec:b5:fa": "Philips Hue",
    # Bilgisayar / yazıcı
    "00:15:60": "HP", "00:1b:78": "HP", "00:21:5a": "HP", "00:23:7d": "HP", "70:5a:0f": "HP",
    "b0:5a:da": "HP",
    "00:12:3f": "Dell", "00:14:22": "Dell", "00:18:8b": "Dell", "00:1c:23": "Dell",
    "b8:2a:72": "Dell", "d4:be:d9": "Dell", "f8:bc:12": "Dell",
    "00:12:fe": "Lenovo", "00:21:cc": "Lenovo", "5c:26:0a": "Lenovo", "e8:6a:64": "Lenovo",
    "00:1f:c6": "ASUS", "00:22:15": "ASUS", "2c:56:dc": "ASUS", "ac:9e:17": "ASUS",
    "04:d4:c4": "ASUS", "1c:87:2c": "ASUS", "50:46:5d": "ASUS",
    "00:00:85": "Canon", "00:1e:8f": "Canon", "2c:9e:fc": "Canon",
    "00:00:48": "Epson", "00:26:ab": "Epson", "a4:ee:57": "Epson", "dc:cc:2d": "Epson",
    "00:80:77": "Brother", "30:05:5c": "Brother", "00:1b:a9": "Brother",
}


def _mac_uretici(mac):
    """OUI tablosundan üretici tahmini; yerel/rastgele MAC'ler ayrı etiketlenir."""
    m = (mac or "").lower().replace("-", ":")
    if len(m) < 8:
        return "bilinmiyor"
    try:
        ilk = int(m.split(":")[0], 16)
    except Exception:
        return "bilinmiyor"
    if ilk & 0x02:
        return "rastgele (gizlilik MAC'i)"
    return MAC_URETICI.get(m[:8], "bilinmiyor")


def _arp_listesi():
    """Windows `arp -a` çıktısını (ip, mac, tur, arayuz_ip) listesine ayıklar."""
    try:
        cikti = subprocess.run(["arp", "-a"], capture_output=True, text=True,
                               timeout=20, errors="replace").stdout or ""
    except Exception:
        return [], ""
    kayitlar, arayuz, sayac = [], "", {}
    for satir in cikti.splitlines():
        a = re.match(r"\s*Interface:\s*(\d+\.\d+\.\d+\.\d+)", satir)
        if a:
            arayuz = a.group(1)
            continue
        m = re.search(r"(\d+\.\d+\.\d+\.\d+)\s+([0-9a-fA-F-]{17})\s+(\w+)", satir)
        if not m:
            continue
        ip, mac2, tur = m.group(1), m.group(2).lower().replace("-", ":"), m.group(3)
        if ip.startswith("224.") or ip.startswith("239.") or mac2.startswith("01:00:5e"):
            continue          # çoklu yayın — cihaz değil
        kayitlar.append({"ip": ip, "mac": mac2, "tur": turel(tur), "arayuz": arayuz})
        sayac[arayuz] = sayac.get(arayuz, 0) + 1
    kendi = ""
    if sayac:
        kendi = max(sayac.items(), key=lambda x: x[1])[0]   # en çok kaydı olan arayüz = yerel ağ
    return kayitlar, kendi


def turel(t):
    return "statik" if str(t).lower().startswith("stat") else "dinamik"


def _ssdp_kesif(sure=3):
    """UDP M-SEARCH (239.255.255.250:1900) ile UPnP cihazlarını bulur."""
    bulunan, yerler = [], set()
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        s.settimeout(0.6)
        msg = ("M-SEARCH * HTTP/1.1\r\nHOST: 239.255.255.250:1900\r\n"
               "MAN: \"ssdp:discover\"\r\nMX: 2\r\nST: ssdp:all\r\n\r\n").encode()
        s.sendto(msg, ("239.255.255.250", 1900))
        bitis = time.time() + max(1, int(sure))
        while time.time() < bitis:
            try:
                ham, adres = s.recvfrom(65535)
            except socket.timeout:
                continue
            except Exception:
                break
            try:
                metin = ham.decode("utf-8", "replace")
            except Exception:
                continue
            m = re.search(r"(?i)^LOCATION:\s*(\S+)", metin, re.M)
            if m:
                yerler.add((adres[0], m.group(1).strip()))
        s.close()
    except Exception:
        return []
    for ip, url in list(yerler)[:25]:
        ad, model = "", ""
        kod, xml, _ = _cek(url, 2)                       # 2 sn zaman aşımı
        if kod == 200 and xml:
            f = re.search(r"<friendlyName>(.*?)</friendlyName>", xml, re.S | re.I)
            mo = re.search(r"<modelName>(.*?)</modelName>", xml, re.S | re.I)
            ad = (f.group(1).strip() if f else "")
            model = (mo.group(1).strip() if mo else "")
        bulunan.append({"ip": ip, "ad": ad, "model": model})
    return bulunan


def yerelag(zorla=False):
    """ARP tablosu + SSDP/UPnP keşfi + MAC üretici tahmini; hafıza veri/yerelag.json."""
    arp, kendi = _arp_listesi()
    ssdp = _ssdp_kesif(3)
    ssdp_yer = {}
    for s in ssdp:
        if s.get("ad") or s.get("model"):
            ssdp_yer.setdefault(s["ip"], s.get("ad") or s.get("model"))

    # hafıza oku (dosya yoksa ilk tarama → hepsi yeni=false olarak kaydedilir)
    hafiza = {}
    ilk_tarama = True
    try:
        with open(YERELAG_HAFIZA, "r", encoding="utf-8") as f:
            hafiza = json.load(f) or {}
        if isinstance(hafiza.get("maclar"), dict):
            ilk_tarama = False
    except Exception:
        hafiza = {}
    bilinen = hafiza.get("maclar") or {}

    cihazlar, yeni_maclar = [], []
    gorulen = set()
    for k in arp:
        if k["mac"] in gorulen:
            continue
        gorulen.add(k["mac"])
        ad = ssdp_yer.get(k["ip"]) or ""
        cihazlar.append({
            "ip": k["ip"], "mac": k["mac"], "uretici": _mac_uretici(k["mac"]),
            "ad": ad, "tur": k["tur"], "yeni": False,
        })
    # NOT: karşılaştırma tabanı döngü içinde DEĞİŞTİRİLMEZ,
    # aksi halde tek taramada 2. cihazdan itibaren hepsi 'yeni' görünür.
    onceki_maclar = set(bilinen.keys())
    for c in cihazlar:
        yeni = bool(onceki_maclar) and c["mac"] not in onceki_maclar   # ilk taramada hepsi eski sayılır
        c["yeni"] = yeni
        if yeni:
            yeni_maclar.append(c["mac"])
        bilinen[c["mac"]] = {"son_gorulme": time.strftime("%Y-%m-%d %H:%M:%S"),
                             "ip": c["ip"], "uretici": c["uretici"]}
    try:
        os.makedirs(VERI, exist_ok=True)
        with open(YERELAG_HAFIZA, "w", encoding="utf-8") as f:
            json.dump({"maclar": bilinen, "guncelleme": time.strftime("%Y-%m-%d %H:%M:%S")},
                      f, ensure_ascii=False, indent=1)
    except Exception:
        pass
    return {
        "cihazlar": cihazlar,
        "kendi_ip": kendi,
        "yeni_cihazlar": yeni_maclar,
        "ssdp": ssdp,
        "cihaz_sayisi": len(cihazlar),
        "not": "ARP tablosu + UPnP keşfi (3 sn)",
        "kaynak": "kendi bilgisayarınız",
        "guncelleme": time.strftime("%Y-%m-%d %H:%M:%S"),
        "ilk_tarama": ilk_tarama,
        "hafiza": YERELAG_HAFIZA,
    }


# =============================================================== 10) ŞİFRE KONTROL
def sifrekontrol(sifre=""):
    """HIBP Pwned Passwords k-anonymity — şifre cihazdan ÇIKMAZ."""
    sifre = str(sifre or "")
    if not sifre:
        return {"hata": "şifre boş olamaz (gövdede {'sifre':'...'} gönderin)"}
    # 1) SHA-1 SADECE YERELDE hesaplanır
    sha1 = hashlib.sha1(sifre.encode("utf-8")).hexdigest().upper()
    ilk5, kalan = sha1[:5], sha1[5:]
    sizinti, bulundu = 0, False
    kod, govde = _cek("https://api.pwnedpasswords.com/range/" + ilk5, 20,
                      {"Add-Padding": "true"})
    if kod == 200 and govde:
        for satir in govde.splitlines():
            p = satir.strip().split(":")
            if len(p) == 2 and p[0].upper() == kalan:
                try:
                    sizinti = int(p[1])
                except Exception:
                    sizinti = 0
                bulundu = sizinti > 0
                break
    else:
        return {"hata": "HIBP sorgusu başarısız (durum %s). Ağ/engel olabilir." % kod}

    # 2) Güç analizi (tamamen yerel)
    uzunluk = len(sifre)
    kumeler = 0
    if re.search(r"[a-z]", sifre):
        kumeler += 26
    if re.search(r"[A-Z]", sifre):
        kumeler += 26
    if re.search(r"\d", sifre):
        kumeler += 10
    if re.search(r"[^A-Za-z0-9]", sifre):
        kumeler += 33
    entropi_kar = _entropi(sifre)
    entropi_bit = round(uzunluk * math.log2(kumeler), 1) if kumeler else 0.0
    puan = entropi_bit
    if bulundu:
        puan = min(puan, 20)               # sızıntıda güç ne olursa olsun zayıftır
    if puan < 40:
        seviye = "zayıf"
    elif puan < 60:
        seviye = "orta"
    elif puan < 80:
        seviye = "güçlü"
    else:
        seviye = "çok güçlü"
    oneriler = []
    if bulundu:
        oneriler.append("Bu şifre %d kez sızmış — DERHAL değiştirin." % sizinti)
    if uzunluk < 12:
        oneriler.append("En az 12 karakter kullanın (16+ ideal).")
    if kumeler <= 26:
        oneriler.append("Büyük/küçük harf, rakam ve simge çeşitliliği ekleyin.")
    if entropi_kar < 3.0:
        oneriler.append("Tekrar eden/öngörülebilir kalıplardan kaçının.")
    if not oneriler:
        oneriler.append("Şifre güçlü görünüyor; her sitede farklı şifre kullanın ve 2FA açın.")
    return {
        "sizinti_sayisi": sizinti,
        "bulundu": bulundu,
        "gucluluk": {"uzunluk": uzunluk, "entropi_bit": entropi_bit,
                     "entropi_karekter": round(entropi_kar, 2),
                     "karakter_havuzu": kumeler, "seviye": seviye},
        "oneriler": oneriler,
        "kaynak": "HIBP Pwned Passwords (k-anonymity)",
        "not": "şifreniz cihazdan çıkmadı; yalnız SHA-1'in ilk 5 karakteri gönderildi",
    }


# ================================================================ 11) RAPOR PAKETİ
def _shading(hucre, renk):
    """python-docx tablo hücresine arka plan rengi (gölge) verir."""
    try:
        from docx.oxml.ns import nsdecls, qn
        from docx.oxml import parse_xml
        hucre._tc.get_or_add_tcPr().append(
            parse_xml(r'<w:shd {} w:val="clear" w:color="auto" w:fill="{}"/>'.format(nsdecls("w"), renk)))
    except Exception:
        pass


def _baslik_seridi(doc, metin, renk="E87722", boyut=13):
    """Renkli, beyaz kalın yazılı başlık şeridi (1 hücreli tablo)."""
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    h = t.rows[0].cells[0]
    _shading(h, renk)
    p = h.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(metin)
    r.bold = True
    r.font.size = Pt(boyut)
    r.font.color.rgb = _rgb("FFFFFF")
    return t


def _rgb(hexs):
    from docx.shared import RGBColor
    return RGBColor.from_string(hexs)


def _tablo(doc, basliklar, satirlar, baslik_rengi="E87722"):
    """Başlık şeritli veri tablosu."""
    from docx.shared import Pt
    t = doc.add_table(rows=1, cols=len(basliklar))
    t.style = "Table Grid"
    for i, b in enumerate(basliklar):
        h = t.rows[0].cells[i]
        _shading(h, baslik_rengi)
        r = h.paragraphs[0].add_run(str(b))
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = _rgb("FFFFFF")
    for satir in satirlar:
        hucreler = t.add_row().cells
        for i, v in enumerate(satir):
            if i < len(hucreler):
                r = hucreler[i].paragraphs[0].add_run(str(v))
                r.font.size = Pt(10)
    return t


def _qr_uret(metin, hedef_png):
    """qrcode → PNG (PIL). Kütüphane yoksa None döner (çökme yok)."""
    try:
        import qrcode
        img = qrcode.make(metin, box_size=6, border=2)
        img.save(hedef_png)
        return hedef_png
    except Exception:
        return None


def raporpaket(govde):
    """Word (.docx) ya da yazdırılabilir tek dosya HTML rapor üretir."""
    govde = govde if isinstance(govde, dict) else {}
    tur = str(govde.get("tur") or "word").lower()
    tur = "html" if tur == "html" else "word"
    baslik = str(govde.get("baslik") or "ÜSTAD OSINT Raporu")
    ozet = govde.get("ozet") if isinstance(govde.get("ozet"), dict) else {}
    bolumler = govde.get("bolumler") if isinstance(govde.get("bolumler"), list) else []
    try:
        os.makedirs(RAPOR_KLASOR, exist_ok=True)
    except Exception as e:
        return {"hata": "rapor klasörü oluşturulamadı: %s (%s)" % (RAPOR_KLASOR, e)}
    zaman = time.strftime("%Y-%m-%d %H:%M:%S")
    ad = time.strftime("%Y-%m-%d-%H%M%S") + "-rapor." + ("docx" if tur == "word" else "html")
    tam = os.path.join(RAPOR_KLASOR, ad)

    # Adli kanıt bloğu: içerik SHA-256 (QR bunu taşır)
    kanit_metni = json.dumps({"baslik": baslik, "ozet": ozet, "bolumler": bolumler, "zaman": zaman},
                             ensure_ascii=False, sort_keys=True)
    kanit_hash = hashlib.sha256(kanit_metni.encode("utf-8")).hexdigest()
    qr_yolu = None

    try:
        if tur == "html":
            govde_html = _html_rapor(baslik, ozet, bolumler, kanit_hash, zaman)
            with open(tam, "w", encoding="utf-8") as f:
                f.write(govde_html)
        else:
            from docx import Document
            from docx.shared import Pt, Cm
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            doc = Document()
            for b in doc.sections:
                b.top_margin = b.bottom_margin = Cm(1.6)
                b.left_margin = b.right_margin = Cm(1.8)
            # --- Ana başlık şeridi (turuncu)
            _baslik_seridi(doc, baslik, "E87722", 15)
            p = doc.add_paragraph()
            r = p.add_run("ÜSTAD OSINT · adli bilişim raporu · " + zaman)
            r.italic = True
            r.font.size = Pt(9)
            # --- Özet tablosu (altın şerit)
            if ozet:
                _baslik_seridi(doc, "ÖZET", "B8860B", 12)
                _tablo(doc, ["Alan", "Değer"], [[str(k), str(v)] for k, v in ozet.items()])
                doc.add_paragraph()
            # --- Bölümler
            for bol in bolumler:
                if not isinstance(bol, dict):
                    continue
                bbaslik = str(bol.get("ad") or "Bölüm")
                _baslik_seridi(doc, bbaslik, "E87722", 12)
                satirlar = bol.get("satirlar") or []
                _tablo(doc, ["#", "Kayıt"],
                       [[i + 1, (str(s)[:800])] for i, s in enumerate(satirlar)])
                doc.add_paragraph()
            # --- Adli kanıt bloğu
            _baslik_seridi(doc, "ADLİ KANIT", "1F3864", 12)
            qr_png = os.path.join(RAPOR_KLASOR, ad.replace(".docx", "") + "-qr.png")
            qr_yolu = _qr_uret("USTAD-OSINT|%s|%s" % (kanit_hash, zaman), qr_png)
            tp = doc.add_paragraph()
            tp.add_run("İçerik SHA-256: ").bold = True
            tp.add_run(kanit_hash)
            tp2 = doc.add_paragraph()
            tp2.add_run("Zaman damgası: ").bold = True
            tp2.add_run(zaman)
            if qr_yolu and os.path.exists(qr_yolu):
                try:
                    doc.add_picture(qr_yolu, width=Cm(3.6))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
                except Exception:
                    qr_yolu = None
            else:
                qr_yolu = None
            # --- Alt bilgi
            alt = doc.add_paragraph()
            ar = alt.add_run("ÜSTAD OSINT · Kenan Kuzucu")
            ar.font.size = Pt(8)
            alt.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.save(tam)
    except Exception as e:
        return {"hata": "rapor üretilemedi (%s): %s" % (tur, e)}

    try:
        boyut = os.path.getsize(tam)
        with open(tam, "rb") as f:
            dosya_hash = hashlib.sha256(f.read()).hexdigest()
    except Exception as e:
        return {"hata": "rapor yazıldı ama okunamadı: %s" % e}
    return {
        "dosya": os.path.abspath(tam), "tur": tur, "boyut": boyut, "sha256": dosya_hash,
        "kanit_sha256": kanit_hash,
        "qr": os.path.abspath(qr_yolu) if (qr_yolu and os.path.exists(qr_yolu)) else None,
        "bolum": len(bolumler), "baslik": baslik,
        "zaman": zaman,
        "kaynak": "yerel üretim (python-docx + qrcode)",
    }


def _html_rapor(baslik, ozet, bolumler, kanit_hash, zaman):
    """Yazdırılabilir tek dosya HTML rapor."""
    def esc(s):
        return _html.escape(str(s))

    o = ['<!DOCTYPE html><html lang="tr"><head><meta charset="utf-8">',
         '<title>%s</title>' % esc(baslik),
         '<style>body{font-family:Segoe UI,Arial,sans-serif;margin:28px;color:#1b1b1b}',
         'h1{background:#E87722;color:#fff;padding:12px 16px;border-radius:6px;font-size:22px}',
         'h2{background:#E87722;color:#fff;padding:8px 12px;border-radius:5px;font-size:16px;margin-top:22px}',
         'h3{background:#1F3864;color:#fff;padding:8px 12px;border-radius:5px;font-size:14px}',
         'table{border-collapse:collapse;width:100%;margin:8px 0 16px}',
         'th{background:#E87722;color:#fff;text-align:left;padding:6px 8px;font-size:13px}',
         'td{border:1px solid #ccc;padding:6px 8px;font-size:12px;vertical-align:top}',
         '.meta{color:#555;font-style:italic;font-size:12px}',
         '.kanit{background:#f2f5fb;border:1px solid #1F3864;padding:10px;border-radius:5px;font-size:12px;word-break:break-all}',
         'footer{text-align:center;color:#666;font-size:11px;margin-top:26px;border-top:1px solid #ddd;padding-top:8px}',
         '@media print{h1,h2,h3{-webkit-print-color-adjust:exact;print-color-adjust:exact}}',
         '</style></head><body>']
    o.append("<h1>%s</h1>" % esc(baslik))
    o.append('<p class="meta">ÜSTAD OSINT · adli bilişim raporu · %s</p>' % esc(zaman))
    if ozet:
        o.append("<h2>ÖZET</h2><table><tr><th>Alan</th><th>Değer</th></tr>")
        for k, v in ozet.items():
            o.append("<tr><td>%s</td><td>%s</td></tr>" % (esc(k), esc(v)))
        o.append("</table>")
    for bol in bolumler:
        if not isinstance(bol, dict):
            continue
        o.append("<h2>%s</h2><table><tr><th style='width:52px'>#</th><th>Kayıt</th></tr>"
                 % esc(bol.get("ad") or "Bölüm"))
        for i, s in enumerate(bol.get("satirlar") or []):
            o.append("<tr><td>%d</td><td>%s</td></tr>" % (i + 1, esc(str(s)[:1200])))
        o.append("</table>")
    o.append("<h3>ADLİ KANIT</h3>")
    o.append('<div class="kanit"><b>İçerik SHA-256:</b> %s<br><b>Zaman damgası:</b> %s</div>'
             % (esc(kanit_hash), esc(zaman)))
    o.append("<footer>ÜSTAD OSINT · Kenan Kuzucu</footer></body></html>")
    return "\n".join(o)


# ==================================================================== 12) BİLDİRİM
BILDIRIM_VARSAYILAN = {"yeni_port": True, "yeni_cve": True, "kritik": True,
                       "whatsapp": True, "eposta": False}


def _bildirim_oku():
    try:
        with open(BILDIRIM_DOSYA, "r", encoding="utf-8") as f:
            d = json.load(f)
        if not isinstance(d, dict):
            raise ValueError
        d.setdefault("ayar", dict(BILDIRIM_VARSAYILAN))
        d.setdefault("olaylar", [])
        for k, v in BILDIRIM_VARSAYILAN.items():
            d["ayar"].setdefault(k, v)
        return d
    except Exception:
        return {"ayar": dict(BILDIRIM_VARSAYILAN), "olaylar": []}


def _bildirim_yaz(d):
    try:
        os.makedirs(VERI, exist_ok=True)
        with open(BILDIRIM_DOSYA, "w", encoding="utf-8") as f:
            json.dump(d, f, ensure_ascii=False, indent=1)
        return True
    except Exception:
        return False


def bildirim(govde):
    """Ayar + olay günlüğü (yerel; ağ kullanmaz)."""
    govde = govde if isinstance(govde, dict) else {}
    d = _bildirim_oku()
    if isinstance(govde.get("ayar"), dict):                 # POST {"ayar": {...}} → kaydet
        for k, v in govde["ayar"].items():
            d["ayar"][str(k)] = bool(v)
        _bildirim_yaz(d)
    if isinstance(govde.get("olay"), dict):                 # POST {"olay": {...}} → ekle
        o = govde["olay"]
        d["olaylar"].insert(0, {
            "tip": str(o.get("tip") or "genel"),
            "baslik": str(o.get("baslik") or "")[:300],
            "onem": str(o.get("onem") or "bilgi"),
            "zaman": time.strftime("%Y-%m-%d %H:%M:%S"),
            "okundu": False,
        })
        d["olaylar"] = d["olaylar"][:200]                   # en fazla 200, yeni başta
        _bildirim_yaz(d)
    olaylar = d.get("olaylar") or []
    return {
        "ayar": d.get("ayar"),
        "olaylar": olaylar[:20],
        "olay_sayisi": len(olaylar),
        "okunmamis": sum(1 for o in olaylar if not o.get("okundu")),
        "kaynak": "yerel olay günlüğü",
    }


def bildirim_olay_ekle(olay):
    """Yalnız olay ekleyip {'tamam':True,'olay_sayisi':N} döner."""
    d = _bildirim_oku()
    o = olay if isinstance(olay, dict) else {}
    d["olaylar"].insert(0, {
        "tip": str(o.get("tip") or "genel"), "baslik": str(o.get("baslik") or "")[:300],
        "onem": str(o.get("onem") or "bilgi"), "zaman": time.strftime("%Y-%m-%d %H:%M:%S"),
        "okundu": False,
    })
    d["olaylar"] = d["olaylar"][:200]
    _bildirim_yaz(d)
    return {"tamam": True, "olay_sayisi": len(d["olaylar"])}


# ==================================================================== YÖNLENDİRİCİ
def ek2c_isle(yol, sorgu, govde, ayarlar=None):
    """"/api/arac/<ad>" yollarını işler. Her modül hatada {"hata": ...} döner, çökmez."""
    ad = str(yol or "").rstrip("/").split("/")[-1].lower()
    sorgu = sorgu if isinstance(sorgu, dict) else {}
    govde = govde if isinstance(govde, dict) else {}
    try:
        if ad in ("moduller", "modul", "ek2c"):
            return {"surum": SURUM, "ekler": [
                {"kod": "saldiriakisi", "ad": "Saldırı Akışı", "aciklama": "SANS ISC en çok saldıran IP'ler + konum"},
                {"kod": "siravc", "ad": "Sır Avcısı", "aciklama": "Yerel dosyalarda maskeli sır/anahtar taraması"},
                {"kod": "yerelag", "ad": "Yerel Ağ", "aciklama": "ARP tablosu + UPnP keşfi + MAC üretici"},
                {"kod": "sifrekontrol", "ad": "Şifre Kontrol", "aciklama": "HIBP k-anonymity sızıntı + güç analizi"},
                {"kod": "raporpaket", "ad": "Rapor Paketi", "aciklama": "Word/HTML adli bilişim raporu"},
                {"kod": "bildirim", "ad": "Bildirim Merkezi", "aciklama": "Uyarı ayarları + olay günlüğü"},
            ]}
        if ad == "saldiriakisi":
            return saldiriakisi(_zorla_mi(sorgu, govde))
        if ad == "siravc":
            return siravc(govde.get("yol") or sorgu.get("yol") or "",
                          govde.get("limit") or sorgu.get("limit") or 400,
                          str(govde.get("github") or sorgu.get("github") or "").lower() in ("1", "true", "evet"),
                          _zorla_mi(sorgu, govde))
        if ad == "yerelag":
            return yerelag(_zorla_mi(sorgu, govde))
        if ad == "sifrekontrol":
            return sifrekontrol(govde.get("sifre") or sorgu.get("sifre") or "")
        if ad == "raporpaket":
            return raporpaket(govde or sorgu)
        if ad == "bildirim":
            if isinstance(govde.get("olay"), dict) and not govde.get("ayar"):
                return bildirim_olay_ekle(govde["olay"])
            return bildirim(govde)
        return {"hata": "bilinmeyen ek2c uç nokta: " + ad, "mevcut": [
            "saldiriakisi", "siravc", "yerelag", "sifrekontrol", "raporpaket", "bildirim", "moduller"]}
    except Exception as e:
        return {"hata": "işlem sırasında beklenmeyen hata: %s: %s" % (type(e).__name__, e)}


# ------------------------------------------------------------------ KENDİ TESTİ
if __name__ == "__main__":
    import sys
    eylem = (sys.argv[1] if len(sys.argv) > 1 else "moduller").lower()
    if eylem == "saldiriakisi":
        r = ek2c_isle("/api/arac/saldiriakisi", {}, None, {})
        print(json.dumps(r, ensure_ascii=False, indent=1)[:1200])
    elif eylem == "yerelag":
        r = ek2c_isle("/api/arac/yerelag", {}, None, {})
        print(json.dumps(r, ensure_ascii=False, indent=1)[:1500])
    elif eylem == "siravc":
        r = ek2c_isle("/api/arac/siravc", {}, {"yol": sys.argv[2] if len(sys.argv) > 2 else KOK,
                                               "limit": 300}, {})
        print(json.dumps(r, ensure_ascii=False, indent=1)[:1200])
    elif eylem == "sifrekontrol":
        print(json.dumps(sifrekontrol(sys.argv[2] if len(sys.argv) > 2 else "123456"),
                         ensure_ascii=False, indent=1)[:800])
    elif eylem == "bildirim":
        print(json.dumps(bildirim({}), ensure_ascii=False, indent=1)[:800])
    elif eylem == "raporpaket":
        print(json.dumps(raporpaket({"tur": "word", "baslik": "Test",
                                     "ozet": {"hedef": "-"}, "bolumler": [{"ad": "B", "satirlar": ["a"]}]}),
                         ensure_ascii=False, indent=1))
    else:
        print(json.dumps(ek2c_isle("/api/arac/moduller", {}, None, {}), ensure_ascii=False, indent=1))
